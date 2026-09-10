#!/usr/bin/env python3
"""Lossless local workbook ingestion with immutable provenance and repeat-safe migration.
Requires openpyxl for reading only. Never modifies the source workbook.
By default creates a review plan; --apply writes it in one SQLite transaction.
"""
import argparse,collections,datetime,decimal,hashlib,json,pathlib,re,sqlite3,zipfile,xml.etree.ElementTree as ET
import openpyxl
D=decimal.Decimal
SCHEMA=[
'CREATE TABLE IF NOT EXISTS workbook_imports(id TEXT PRIMARY KEY,tenant_id TEXT NOT NULL,filename TEXT NOT NULL,sha256 TEXT NOT NULL,metadata TEXT NOT NULL,original BLOB NOT NULL,created TEXT NOT NULL,UNIQUE(tenant_id,sha256))',
'CREATE TABLE IF NOT EXISTS workbook_sheets(batch_id TEXT NOT NULL,sheet_index INTEGER NOT NULL,name TEXT NOT NULL,metadata TEXT NOT NULL,PRIMARY KEY(batch_id,sheet_index))',
'CREATE TABLE IF NOT EXISTS workbook_rows(batch_id TEXT NOT NULL,sheet_index INTEGER NOT NULL,row_number INTEGER NOT NULL,data TEXT NOT NULL,PRIMARY KEY(batch_id,sheet_index,row_number))',
'CREATE TABLE IF NOT EXISTS workbook_facts(id TEXT PRIMARY KEY,batch_id TEXT NOT NULL,tenant_id TEXT NOT NULL,domain TEXT NOT NULL,data TEXT NOT NULL)',
'CREATE INDEX IF NOT EXISTS workbook_facts_scope ON workbook_facts(tenant_id,batch_id,domain)']
def encode(x):return json.dumps(x,ensure_ascii=False,default=lambda v:v.isoformat() if hasattr(v,'isoformat') else str(v),separators=(',',':'))
def text(v):return '' if v is None else re.sub(r'\s+',' ',str(v)).strip()
def number(v):
 if v is None or isinstance(v,bool) or text(v) in ('','-','—'):return None
 try:
  n=D(str(v));return n if n.is_finite() else None
 except decimal.InvalidOperation:return None
def cents(v):return int((number(v) or D(0)).quantize(D('.01'),rounding=decimal.ROUND_HALF_UP)*100)
def date(v):
 if isinstance(v,(datetime.date,datetime.datetime)):return v.isoformat()[:10] if v.year>=1950 else None
 if isinstance(v,str) and re.match(r'^\d{4}-\d{2}-\d{2}',v):return v[:10] if v[:4]>='1950' else None
 return None
def fy(v):
 if not date(v):return 'undated'
 y=int(date(v)[:4])-(int(date(v)[5:7])<4);return f'{y}–{str(y+1)[-2:]}'
def run(args):
 source=pathlib.Path(args.workbook);raw=source.read_bytes();digest=hashlib.sha256(raw).hexdigest();batch='wb-'+hashlib.sha256((args.tenant+digest).encode()).hexdigest()[:24]
 output=pathlib.Path(args.output);output.mkdir(parents=True,exist_ok=True)
 w=openpyxl.load_workbook(source,data_only=False);cached=openpyxl.load_workbook(source,data_only=True)
 sheets=[];archived=[];facts=[];issues=[];source_rows={};records={};now=datetime.datetime.now(datetime.timezone.utc).isoformat()
 def uid(kind,key):return 'wb-'+hashlib.sha256((batch+kind+str(key)).encode()).hexdigest()[:28]
 def issue(code,sheet,row,detail,level='Review'):issues.append({'code':code,'sheet':sheet,'row':row,'detail':detail,'level':level})
 def put(kind,key,data,origin=None,master=False):
  i=uid(kind,key);data={'operation':True,'status':'Active' if master else 'Imported','importBatch':batch,'importLocked':not master,**data}
  if origin:data['sourceWorkbook']={'batchId':batch,'sheet':origin[0],'rows':origin[1] if isinstance(origin[1],list) else [origin[1]]}
  records[i]={'id':i,'kind':kind,'fy':'master' if master else fy(data.get('date')),'data':data};return i
 for si,s in enumerate(w):
  by=collections.defaultdict(list);rc=collections.defaultdict(dict);formula_count=0
  for row in s:
   for c in row:
    if c.value is None and not c.comment and not c.hyperlink:continue
    v=c.value.text if hasattr(c.value,'text') else c.value;cv=cached[s.title][c.coordinate].value
    cell={'address':c.coordinate,'column':c.column,'value':v,'cached':cv,'type':c.data_type,'numberFormat':c.number_format,'comment':c.comment.text if c.comment else None,'hyperlink':c.hyperlink.target if c.hyperlink else None,'arrayRef':getattr(c.value,'ref',None)}
    by[c.row].append(cell);rc[c.row][c.column_letter if hasattr(c,'column_letter') else openpyxl.utils.get_column_letter(c.column)]=cv
    if c.data_type=='f':formula_count+=1
    if cached[s.title][c.coordinate].data_type=='e':issue('SOURCE_FORMULA_ERROR',s.title,c.row,f'{c.coordinate}: {cv}')
  source_rows[s.title]=rc
  meta={'dimension':s.calculate_dimension(),'maxRow':s.max_row,'maxColumn':s.max_column,'state':s.sheet_state,'cellCount':sum(map(len,by.values())),'formulaCount':formula_count,'tables':[{ 'name':t.name,'ref':t.ref,'headers':[x.name for x in t.tableColumns]} for t in s.tables.values()],'mergedRanges':[str(x) for x in s.merged_cells.ranges],'hiddenRows':[n for n,d in s.row_dimensions.items() if d.hidden],'hiddenColumns':[n for n,d in s.column_dimensions.items() if d.hidden],'validations':[str(x) for x in s.data_validations.dataValidation]}
  sheets.append({'index':si,'name':s.title,**meta})
  archived.extend((si,r,encode(cells)) for r,cells in by.items())
  if s.title not in ['Raw Data_Sales','Raw Data_Purchase','Expenses','Raw Data (FY 2025-26)']:continue
  for r,v in rc.items():
   old=s.title=='Raw Data (FY 2025-26)';domain='sales' if s.title in ['Raw Data_Sales','Raw Data (FY 2025-26)'] else 'purchase' if s.title=='Raw Data_Purchase' else 'expense'
   if r<=(2 if old else 1):continue
   typ=text(v.get('C' if old else 'B'));party=text(v.get('D' if domain=='sales' else 'C'))
   if not typ and not party:continue
   f={'id':uid('fact',f'{si}:{r}'),'sheet':s.title,'sheetIndex':si,'row':r,'domain':domain,'type':typ,'party':party,'invoice':text(v.get('E' if domain=='sales' else 'D')),'date':date(v.get('F' if domain=='sales' else 'E')),'priorYear':old}
   mappings=(dict(qty='H',description='G',fob='K',freight='L',insurance='M',usd='N',exchangeRate='O',basic='P',gstRate='Q',sgst='R',cgst='S',igst='T',invoiceValue='W' if old else 'U',payable='AB' if old else 'W',tcs='V',cashUsd='AC' if old else 'Y',charges='Z' if not old else 'ZZ',settledUsd='AA' if not old else 'AC',conversionRate='AD' if old else 'AB',paid='AE' if old else 'AC',paymentDate='AF' if old else 'AD',bankRef='AG' if old else 'AE',workflow='C' if not old else 'ZZ',incoterm='I',hsn='J',shippingBill='AR' if old else 'AM',shippingDate='AS' if old else 'AN',leo='AO' if not old else 'ZZ',port='BA' if old else 'AQ',country='AQ' if old else 'AR',dischargePort='AS' if not old else 'ZZ',vessel='AT' if not old else 'ZZ',etd='AO' if old else 'AU',eta='AP' if old else 'AV',bl='AT' if old else 'AW',blDate='AU' if old else 'AX',egm='AX' if old else 'BE',egmDate='AY' if old else 'BF',ebrc='BC' if old else 'BH',ebrcDate='BD' if old else 'BI',dbk='BI' if old else 'BO',dbkReceived='BN' if old else 'BR',dbkDate='BO' if old else 'BS',rodtep='BM' if old else 'BQ',igstReceived='BW' if old else 'BU',igstDate='BX' if old else 'BV',fx='BK' if old else 'BM') if domain=='sales' else dict(qty='M',serial='I',machineType='J',model='K',brand='L',hsn='H',gstin='F',pan='G',basic='O',gstRate='P',sgst='Q',cgst='R',igst='S',tcs='U',other='V',invoiceValue='W',tdsRate='X',tds='Y',payable='Z',paid='AB',paymentDate='AC',bankRef='AD',inventoryStatus='AP',linkedInvoice='AQ',sourceSalesDate='AR',tdsDeductionDue='AW',tdsDeducted='AX',tdsDepositDue='BA',tdsDeposited='BB',tdsDeductionStatus='AY',tdsDepositStatus='BC',tdsDeductionInterest='AZ',tdsDepositInterest='BD',tdsSection='AU',serialValidation='BN',shipFrom='AL',shipTo='AM',transporter='AN',eway='AF',ewayDate='AG',vehicle='AH') if domain=='purchase' else dict(linkedInvoice='F',gstin='G',pan='H',basic='I',sgst='J',cgst='K',igst='L',invoiceValue='M',tdsRate='N',tds='O',payable='P',paid='Q',paymentDate='R',bankRef='S'))
   for key,col in mappings.items():
    val=v.get(col)
    f[key]=date(val) if key.endswith('Date') or key in ['tdsDeductionDue','tdsDeducted','tdsDepositDue','tdsDeposited','etd','eta'] else (str(val) if number(val) is not None else text(val))
   f['fy']=fy(f['date']);f['hasPayment']=bool(number(f.get('paid')) or number(f.get('cashUsd')))
   if domain=='expense' and number(f['payable']) is None:
    f['payable']=str((number(f['invoiceValue']) or D(0))-(number(f['tds']) or D(0)));f['payableDerived']=True
   if old and number(f['payable']) is None:f['payable']=f['invoiceValue'];f['payableDerived']=True
   if not f['date'] and typ not in ['Advance','Advance Receipts']:issue('MISSING_DATE',s.title,r,'Transaction date is absent; no date invented.')
   if domain=='purchase' and typ=='Purchase' and text(f.get('linkedInvoice')) and text(f.get('sourceSalesDate')).startswith('#'):issue('UNRESOLVED_SALES_LOOKUP',s.title,r,'The source lookup failed. Match sales, prior-year sales and purchase references separately.')
   if domain=='expense' and ',' in f['linkedInvoice']:issue('UNALLOCATED_SHARED_EXPENSE',s.title,r,'Multiple invoice references without amounts per invoice. Preserved as one shared expense, not duplicated.')
   if domain=='purchase' and f.get('serialValidation','').startswith('⚠'):issue('SERIAL_WARNING',s.title,r,f['serialValidation'])
   if f['hasPayment'] and not f.get('paymentDate'):issue('MISSING_PAYMENT_DATE',s.title,r,'Payment amount exists without a usable date.')
   for key in ['basic','payable','paid']:
    n=number(f.get(key))
    if n is not None and abs(n-n.quantize(D('.01'),rounding=decimal.ROUND_HALF_UP))>D('.000001'):issue('SUB_PAISE_PRECISION',s.title,r,f'{key}: source precision retained; ERP monetary display rounds to paise.','Information')
   facts.append(f)
 # Source controls deliberately exclude template rows and total rows.
 controls={}
 for domain in ['sales','purchase','expense']:
  for prior in ([False,True] if domain=='sales' else [False]):
   group=[f for f in facts if f['domain']==domain and f['priorYear']==prior]
   controls[domain+('-prior' if prior else '')]={'rows':len(group),**{k:str(sum((number(f.get(k)) or D(0) for f in group),D(0))) for k in ['basic','invoiceValue','payable','paid','usd','cashUsd','settledUsd','dbkReceived','igstReceived']}}
 # Invoice/detail mappings. One real invoice may contain multiple machines or allocations.
 currencies={c:put('master-currency',c,{'name':c,'code':c,'words':c},master=True) for c in ['INR','USD']}
 unit=put('master-units','NOS',{'name':'Numbers','code':'NOS'},master=True)
 parties={};products={};invoice_lookup=collections.defaultdict(list)
 def party(f):
  k='customers' if f['domain']=='sales' else 'vendors';key=(k,text(f['party']).casefold())
  if not key[1]:return None
  if key not in parties:parties[key]=put(k,key,{'name':f['party'],'gstin':f.get('gstin',''),'pan':f.get('pan','')},(f['sheet'],f['row']),True)
  else:
   p=records[parties[key]]['data']
   if f.get('gstin') and p.get('gstin') and f['gstin']!=p['gstin']:issue('PARTY_GST_CONFLICT',f['sheet'],f['row'],'Same party name has different GST registrations; original registrations retained in source rows.')
   if not p.get('gstin') and f.get('gstin'):p['gstin']=f['gstin'];p['pan']=f.get('pan','')
  return parties[key]
 for f in facts:f['partnerId']=party(f)
 for f in facts:
  if f['domain']=='purchase' and f['type']=='Purchase':
   key='|'.join(text(f.get(x)) for x in ['brand','machineType','model','hsn'])
   if key not in products:products[key]=put('products',key,{'name':' '.join(text(f.get(x)) for x in ['brand','machineType','model']).strip() or 'Purchase item (description not supplied)','sku':'WB-'+hashlib.sha256(key.encode()).hexdigest()[:10].upper(),'unitId':unit,'hsn':f['hsn'],'brand':f['brand'],'model':f['model'],'machineType':f['machineType']},(f['sheet'],f['row']),True)
   f['productId']=products[key]
 # Populate configuration lists explicitly present in the workbook.
 expense_types={}
 for rn,values in source_rows.get('Validations',{}).items():
  label=text(values.get('M'))
  if label:expense_types[label]=put('master-expense-types',label,{'name':label},('Validations',rn),True)
 for f in facts:
  if f['domain']=='sales':
   for field,kind in [('incoterm','master-shipment-terms'),('port','master-ports')]:
    value=text(f.get(field))
    if value:put(kind,value,{'name':value,'code':value},(f['sheet'],f['row']),True)
 groups=collections.defaultdict(list)
 for f in facts:
  if f['domain']=='sales':
   if f['type'] in ['Part Receipts','Advance Receipts'] or not number(f['invoiceValue']):continue
   kind='sales-credit-notes' if f['type']=='Credit Note' else 'invoices' if 'Export' in f['type'] else 'domestic-invoices'
  elif f['domain']=='purchase':
   if f['type'] not in ['Purchase','Debit Note']:continue
   kind='purchase-debit-notes' if f['type']=='Debit Note' else 'purchase-invoices'
  else:kind='expense-credit-notes' if f['type']=='Credit Note' else 'expenses'
  key=(kind,f['partnerId'],f['invoice'] or 'MISSING-'+str(f['row']),f['fy'])
  groups[key].append(f)
 refs=collections.Counter((k[0],k[2]) for k in groups)
 for key,fs in groups.items():
  kind,partner,reference,_=key;f=fs[0];outref=reference if refs[(kind,reference)]==1 else reference+' · '+hashlib.sha256(str(partner).encode()).hexdigest()[:6]
  amt=sum((number(x['payable']) or D(0) for x in fs),D(0));export=kind=='invoices'
  lines=[]
  for x in fs:
   lines.append({'description':x.get('description') or ' '.join(text(x.get(k)) for k in ['brand','machineType','model']) or x['type'],'productId':x.get('productId'),'quantity':x.get('qty',''),'uom':'NOS' if number(x.get('qty')) else '', 'rate':'','basic':cents(x.get('basic')),'gst':cents(sum((number(x.get(k)) or D(0) for k in ['sgst','cgst','igst']),D(0))),'tcs':cents(x.get('tcs')),'tds':cents(x.get('tds')),'total':cents(x['payable']),'serialNumbers':x.get('serial',''),'sourceRow':x['row']})
  record={'reference':outref,'originalReference':reference,'date':f['date'],'partnerId':partner,'customerName':f['party'],'currency':'INR','currencyId':currencies['INR'],'amount':cents(amt),'sourceAmountExact':str(amt),'lines':lines,'quotationType':'Export' if export else 'Domestic','sourceTransactionType':f['type'],'sourceStatus':f.get('workflow') or f.get('inventoryStatus'),'status':'Cancelled' if 'Cancelled' in f['type'] else 'Imported','sourceFacts':[x['id'] for x in fs],'totalUsd':str(sum((number(x.get('usd')) or D(0) for x in fs),D(0))),'expenseTypeId':expense_types.get(f['type']),'amountDerived':any(x.get('payableDerived') for x in fs),'sourceInvoiceValue':str(sum((number(x['invoiceValue']) or D(0) for x in fs),D(0))),'notes':'Imported historical source. Not posted to the general ledger.'}
  rid=put(kind,key,record,(f['sheet'],[x['row'] for x in fs]));invoice_lookup[(f['domain'],partner,reference)].append(rid)
  for x in fs:x['recordId']=rid;x['recordKind']=kind
 # Link payments to invoice groups rather than duplicating invoices for payment rows.
 for f in facts:
  candidates=invoice_lookup[(f['domain'],f['partnerId'],f['invoice'])];candidates=[i for i in candidates if records[i]['kind'] not in ['sales-credit-notes','purchase-debit-notes','expense-credit-notes']]
  target=candidates[0] if len(candidates)==1 else None
  if not f.get('recordId') and target:f['recordId']=target;f['recordKind']=records[target]['kind']
  if f['hasPayment']:
   advance=f['type'] in ['Advance','Advance Receipts'];kind=('customer-advances' if advance else 'receipts') if f['domain']=='sales' else ('supplier-advances' if advance else 'payments')
   usd=number(f.get('cashUsd'));currency='USD' if f['domain']=='sales' and usd else 'INR'
   d={'reference':f'WB-{kind.upper()}-{f["sheetIndex"]+1}-{f["row"]}','date':f['paymentDate'],'partnerId':f['partnerId'],'invoiceId':target,'sourceId':target,'currency':currency,'currencyId':currencies[currency],'amount':cents(usd if currency=='USD' else f['paid']),'settledAmount':cents(f.get('settledUsd') if currency=='USD' else f['paid']),'bankChargesUsd':f.get('charges'),'inrReceivedExact':f['paid'],'bankReference':f['bankRef'],'sourceFacts':[f['id']]}
   f['paymentRecordId']=put(kind,f['id'],d,(f['sheet'],f['row']))
 # Explicit links across current and prior years, preserving unmatched/multi-target references.
 sales_refs=collections.defaultdict(list);purchase_refs=collections.defaultdict(list)
 for i,r in records.items():
  if r['kind'] in ['invoices','domestic-invoices']:sales_refs[r['data']['originalReference']].append(i)
  if r['kind']=='purchase-invoices':purchase_refs[r['data']['originalReference']].append(i)
 serials=collections.Counter(text(f.get('serial')) for f in facts if f['domain']=='purchase' and f['type']=='Purchase' and f.get('serial'))
 for f in facts:
  if f['domain'] in ['purchase','expense'] and f.get('linkedInvoice'):
   links=[]
   for token in [text(t) for t in f['linkedInvoice'].split(',')]:links.extend(sales_refs.get(token,[]) or purchase_refs.get(token,[]))
   f['linkedRecordIds']=list(dict.fromkeys(links))
   if len(links)==1 and f.get('recordId'):records[f['recordId']]['data'].setdefault('relatedIds',[]).append(links[0])
   if not links and f.get('linkedInvoice'):issue('UNMATCHED_REFERENCE',f['sheet'],f['row'],'Reference is not a unique invoice in the supplied workbook: '+f['linkedInvoice'])
  if f['domain']=='purchase' and f['type']=='Purchase':
   sn=text(f.get('serial'));is_serial=bool(sn) and '\n' not in f.get('serial','') and sn.casefold() not in ['multiple parts','spare parts'] and ' ' not in sn
   if is_serial and serials[sn]==1:
    state=text(f['inventoryStatus']).lower();status='Sold' if state=='sold' else 'Consumed' if state=='consumed' else 'Available' if state in ['','not sold'] else 'Review'
    f['machineRecordId']=put('machines',sn,{'name':sn,'serialNumber':sn,'productId':f['productId'],'purchaseId':f.get('recordId'),'salesInvoiceId':(f.get('linkedRecordIds') or [None])[0],'cost':f['basic'],'date':f['date'],'status':status,'locationStatus':'Not supplied in workbook','sourceFacts':[f['id']]},(f['sheet'],f['row']))
   elif sn:issue('NON_UNIQUE_OR_NON_SERIAL_ITEM',f['sheet'],f['row'],'Serial description preserved without inventing individual machine identities.')
  if f['domain']=='sales' and f.get('recordId') and f['type'] not in ['Credit Note','Part Receipts','Advance Receipts']:
   parent=f['recordId'];sb=None
   if f.get('shippingBill'):
    sb=put('shipping-bills',f['id'],{'reference':f['shippingBill']+' · '+f['invoice'],'number':f['shippingBill'],'date':f['shippingDate'],'sourceId':parent,'sourceStatus':f.get('leo'),'assessedDate':f['shippingDate'],'egmNumber':f['egm'],'egmDate':f['egmDate'],'portCode':f['port'],'country':f['country'],'sourceFacts':[f['id']]},(f['sheet'],f['row']))
   if f.get('bl'):put('bills-of-lading',f['id'],{'reference':f['bl']+' · '+f['invoice'],'date':f['blDate'],'sourceId':sb or parent,'vessel':f['vessel'],'etd':f['etd'],'eta':f['eta'],'sourceFacts':[f['id']]},(f['sheet'],f['row']))
   if f.get('ebrc'):put('ebrc',f['id'],{'reference':f['ebrc']+' · '+f['invoice'],'number':f['ebrc'],'date':f['ebrcDate'],'sourceId':parent,'sourceFacts':[f['id']]},(f['sheet'],f['row']))
   for scheme,entitlement,received,dt in [('DBK','dbk','dbkReceived','dbkDate'),('RoDTEP','rodtep',None,None),('IGST refund','igst','igstReceived','igstDate')]:
    if number(f.get(entitlement)) or (received and number(f.get(received))):put('incentives',f['id']+scheme,{'reference':f['invoice']+' · '+scheme,'date':f['date'],'sourceId':sb or parent,'invoiceId':parent,'scheme':scheme,'currency':'INR','amount':cents(f.get(entitlement)),'receivedAmount':f.get(received) if received else None,'receivedDate':f.get(dt) if dt else None,'sourceFacts':[f['id']]},(f['sheet'],f['row']))
 # Detect suspicious duplication without silently deleting any source row.
 seen={}
 for f in facts:
  if f['domain']!='expense':continue
  key=(f['party'].casefold(),f['invoice'],f['basic'],f['invoiceValue'])
  if key in seen:issue('REPEATED_EXPENSE_VALUE',f['sheet'],f['row'],f'Same supplier/invoice/value as row {seen[key]}. Separate allocation links are preserved; an empty link may indicate a duplicate.')
  else:seen[key]=f['row']
 for f in facts:
  if f['domain']=='expense' and f.get('linkedInvoice'):
   machine=next((r for r in records.values() if r['kind']=='machines' and r['data']['serialNumber']==f['linkedInvoice']),None)
   if machine and f.get('recordId'):
    records[f['recordId']]['data'].setdefault('relatedIds',[]).append(machine['id']);f['linkedRecordIds']=[machine['id']]
    issues[:]=[i for i in issues if not(i['code']=='UNMATCHED_REFERENCE' and i['sheet']==f['sheet'] and i['row']==f['row'])]
 with zipfile.ZipFile(source) as z:
  for name in z.namelist():
   if 'pivotCacheDefinition' in name and name.endswith('.xml'):
    node=ET.fromstring(z.read(name));src=next((x for x in node.iter() if x.tag.endswith('worksheetSource')),None)
    domain={'tblSales':'sales','tblPurchase':'purchase'}.get(src.get('name') if src is not None else '')
    current=sum(f['domain']==domain and not f['priorYear'] for f in facts)
    if domain and int(node.get('recordCount','0'))<current:issue('STALE_PIVOT_CACHE','Pivot',1,'Saved pivot cache has '+node.get('recordCount','0')+' rows for '+str(domain)+'; source contains '+str(current)+' business rows. ERP monthly reports use all source rows.')
 if 'Gross Profit' in source_rows:
  issue('SOURCE_SUMMARY_DOUBLE_COUNT','Gross Profit',2,'E2 sums the entire drawback receipt column, including its total row. ERP reports sum transaction rows only.')
  issue('SOURCE_SUMMARY_RANGE','Gross Profit',17,'E17 sums only I24:I68 and omits later invoice profit rows. ERP reports include all matched invoice rows.')
  issue('SOURCE_POSITIONAL_STOCK_LINK','Gross Profit',10,'E10 points to a positional Inventory cell. Stock categories must use item attributes, not a sorted row number.')
 for f in facts:
  if f['domain']=='sales' and f['type']=='Part Receipts' and number(f.get('fx')) and not number(f.get('exchangeRate')):issue('FX_BASE_RATE_MISSING',f['sheet'],f['row'],'Receipt-row FX formula uses a blank invoice exchange rate. Original result retained; reliable FX requires the linked invoice rate.')
 metadata={'batchId':batch,'filename':source.name,'sha256':digest,'sheets':sheets,'cellCount':sum(s['cellCount'] for s in sheets),'formulaCount':sum(s['formulaCount'] for s in sheets),'sourceRowCount':len(facts),'controls':controls,'issues':issues,'recordCounts':dict(collections.Counter(r['kind'] for r in records.values())),'created':now,'accountingStatus':'Historical operational import; general ledger not posted','sourceWorkbookUnchanged':True}
 plan={'metadata':metadata,'facts':facts,'records':list(records.values())}
 (output/'import-plan.json').write_text(encode(plan));(output/'reconciliation.json').write_text(encode({'controls':controls,'issues':issues,'recordCounts':metadata['recordCounts']}))
 if args.apply:
  if not args.database:raise ValueError('--database is required')
  db=sqlite3.connect(args.database);db.execute('PRAGMA foreign_keys=ON')
  if not db.execute('SELECT id FROM tenants WHERE id=?',(args.tenant,)).fetchone():raise ValueError('Unknown tenant')
  for sql in SCHEMA:db.execute(sql)
  if db.execute('SELECT id FROM workbook_imports WHERE tenant_id=? AND sha256=?',(args.tenant,digest)).fetchone():print('Already imported; no rows changed.');return
  if db.execute('SELECT id FROM workbook_imports WHERE tenant_id=? AND filename=?',(args.tenant,source.name)).fetchone():raise ValueError('A different revision of this workbook was already imported. Reconcile changes before replacing it.')
  before=dict(db.execute('SELECT id,data FROM records'))
  with db:
   db.execute('INSERT INTO workbook_imports VALUES(?,?,?,?,?,?,?)',(batch,args.tenant,source.name,digest,encode(metadata),raw,now))
   db.executemany('INSERT INTO workbook_sheets VALUES(?,?,?,?)',[(batch,s['index'],s['name'],encode(s)) for s in sheets])
   db.executemany('INSERT INTO workbook_rows VALUES(?,?,?,?)',[(batch,*r) for r in archived])
   db.executemany('INSERT INTO workbook_facts VALUES(?,?,?,?,?)',[(f['id'],batch,args.tenant,f['domain'],encode(f)) for f in facts])
   db.executemany('INSERT INTO records(id,tenant_id,kind,fy,data,created) VALUES(?,?,?,?,?,?)',[(r['id'],args.tenant,r['kind'],r['fy'],encode(r['data']),now) for r in records.values()])
   actor=db.execute("SELECT id FROM users WHERE tenant_id=? AND role='Admin' ORDER BY id LIMIT 1",(args.tenant,)).fetchone()[0]
   db.execute('INSERT INTO audit VALUES(?,?,?,?,?,?,?,?)',(uid('audit','import'),args.tenant,actor,'Workbook imported',batch,'workbook-import',encode({'sourceHash':digest,'sourceRows':len(facts),'recordCounts':metadata['recordCounts']}),now))
   assert all(db.execute('SELECT data FROM records WHERE id=?',(i,)).fetchone()[0]==d for i,d in before.items())
   assert db.execute('SELECT count(*) FROM workbook_facts WHERE batch_id=?',(batch,)).fetchone()[0]==len(facts)
   assert sum(len(json.loads(d)) for d, in db.execute('SELECT data FROM workbook_rows WHERE batch_id=?',(batch,)))==metadata['cellCount']
  assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
  (output/'applied.json').write_text(encode({'batchId':batch,'records':len(records),'cells':metadata['cellCount'],'originalSha256':digest,'existingRecordsUnchanged':len(before)}))
  print('APPLIED',batch,len(records),'records')
 print(encode({k:metadata[k] for k in ['batchId','cellCount','formulaCount','sourceRowCount','recordCounts','controls']}));print('Issues by code:',dict(collections.Counter(i['code'] for i in issues)))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('--tenant',default='initial');p.add_argument('--database');p.add_argument('--output',required=True);p.add_argument('--apply',action='store_true');run(p.parse_args())
